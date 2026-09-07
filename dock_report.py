#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
dock_report.py
==============

Erzeugt aus einer Docking-Kaskade eine uebersichtliche Word-Auswertung.

Ablauf
------
1.  Durchsucht <BASE> nach allen Unterordnern ausser den ausgeschlossenen
    (Default: "_unpacked").  Jeder dieser Ordner ist eine "Pose-Serie"
    (z.B. eS4_7b7d_p0).
2.  Liest aus jedem *.pdbqt nur das erste MODEL und extrahiert dort die
    Zeile(n)  "REMARK SMILES ..."  (Fortsetzungszeilen werden zusammen-
    gehaengt, "REMARK SMILES IDX" wird ignoriert).
3.  Sucht rekursiv unter <BASE>/_unpacked nach  Top250_<ordner>.csv  und
    liest daraus RANK, ECR-Score, score_vina_best,
    score_dense_cnnaffinity_best und Pose.
4.  Verknuepft beides ueber den Molekuel-Namen (bevorzugt) bzw. ueber den
    Rank aus dem Dateinamen (Fallback).
5.  Schreibt einen formatierten Word-Report (+ optional CSV/XLSX).

Aufruf
------
    python3 dock_report.py                          # Defaults benutzen
    python3 dock_report.py --base /pfad/extr_ --out report.docx
    python3 dock_report.py --top 50 --also-csv
    python3 dock_report.py --dump-headers           # nur CSV-Spalten zeigen

Abhaengigkeiten
---------------
    pip install python-docx
    (optional fuer --also-xlsx:  pip install openpyxl)
"""

from __future__ import annotations

import argparse
import csv
import datetime as _dt
import re
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

# ----------------------------------------------------------------------
# Konfiguration (kann per CLI ueberschrieben werden)
# ----------------------------------------------------------------------

DEFAULT_BASE = Path("/itpstore/roessel/project_sbg/kepler-dock-cascade/extr_")
EXCLUDE_DIRS = {"_unpacked"}
UNPACKED_DIRNAME = "_unpacked"
TOP_CSV_GLOB = "Top250_*.csv"
RENAME_MAP_NAME = "rename_map.csv"      # von rename_pdbqt.py erzeugt
SELECTION_CSV_NAME = "selection.csv"    # von extract_top_ligands.py erzeugt
MANIFEST_NAMES = ("MANIFEST.txt", "manifest.txt", "config", "config.ini")

# PDBQT-Dateien, die keine Ligandenpose sind
EXCLUDE_PDBQT = {"00_target.pdbqt"}
EXCLUDE_PDBQT_SUFFIX = ("_target.pdbqt", "_receptor.pdbqt")

# Anzeigenamen der ECR-Gewichte aus dem MANIFEST
WEIGHT_LABELS = {
    "w_vina": "AutoDock Vina",
    "w_vinardo": "Vinardo",
    "w_ad4": "AutoDock4",
    "w_cnnaffinity": "GNINA CNNaffinity (pK)",
    "w_cnnscore": "GNINA CNNscore (0-1)",
    "w_deltalinf9xgb": "Delta-Lin_F9XGB",
    "w_dense_cnnaffinity": "GNINA dense CNNaffinity (pK)",
    "w_dense_cnnscore": "GNINA dense CNNscore (0-1)",
}

# Kandidaten fuer die CSV-Spaltennamen. Erst exakte Treffer (case-insensitiv,
# ohne Sonderzeichen), dann Teilstring-Treffer. Bei abweichenden Headern hier
# einfach den echten Namen vorne einfuegen.
COLUMN_CANDIDATES: Dict[str, Dict[str, Sequence[str]]] = {
    "rank":  {"exact": ("rank", "ranking", "no", "nr", "index", "idx"),
              "contains": ("rank",)},
    "name":  {"exact": ("name", "ligand", "ligand_name", "title", "molname",
                        "mol_name", "id", "molecule", "compound"),
              "contains": ("ligand", "name", "title", "mol")},
    "ecr":   {"exact": ("ecr_score", "ecr", "score_ecr", "ecrscore"),
              "contains": ("ecr",)},
    "vina":  {"exact": ("score_vina_best", "vina_best", "vina"),
              "contains": ("vina",)},
    "cnn":   {"exact": ("score_dense_cnnaffinity_best",
                        "dense_cnnaffinity_best", "cnnaffinity_best"),
              "contains": ("cnnaffinity", "cnn_affinity", "cnnaff")},
    "pose":  {"exact": ("best_pose", "pose", "pose_id", "mode", "model"),
              "contains": ("pose",)},
    "file":  {"exact": ("file", "filename", "pdbqt", "pose_file"),
              "contains": ("filename", "pdbqt")},
    "smiles": {"exact": ("smiles", "canonical_smiles", "smi"),
               "contains": ("smiles",)},
}

# Dateiname:  0001_Drugs_inhibitors_sho_mol_0022884_docked.pdbqt
FNAME_RE = re.compile(
    r"^(?P<rank>\d+)_(?P<name>.+?)(?:_docked)?\.pdbqt$", re.IGNORECASE
)

ZWSP = "\u200b"  # Zero-Width-Space: erlaubt Word den Umbruch langer SMILES


# ----------------------------------------------------------------------
# Hilfsfunktionen
# ----------------------------------------------------------------------

def norm(s: str) -> str:
    """Header/Namen normalisieren: klein, nur alphanumerisch."""
    return re.sub(r"[^a-z0-9]+", "", (s or "").lower())


def find_col(fieldnames: Sequence[str], key: str) -> Optional[str]:
    """Passende CSV-Spalte fuer einen logischen Key finden."""
    cand = COLUMN_CANDIDATES[key]
    normalized = {norm(f): f for f in fieldnames if f}
    for c in cand["exact"]:
        if norm(c) in normalized:
            return normalized[norm(c)]
    for c in cand["contains"]:
        for nf, orig in normalized.items():
            if norm(c) in nf:
                return orig
    return None


def strip_mol_key(value: str) -> str:
    """
    Aus beliebigen Bezeichnern einen stabilen Join-Key machen.

    '0001_Drugs_inhibitors_sho_mol_0022884_docked.pdbqt'
    'Drugs_inhibitors_sho_mol_0022884'
    '/pfad/0001_Drugs_..._docked.pdbqt'
        ->  'drugsinhibitorsshomol0022884'
    """
    v = str(value or "").strip()
    v = Path(v).name
    for suf in (".pdbqt", ".sdf", ".mol2", ".pdb"):
        if v.lower().endswith(suf):
            v = v[: -len(suf)]
    v = re.sub(r"^\d{1,6}[_\-]", "", v)          # fuehrender Rank
    v = re.sub(r"[_\-]docked$", "", v, flags=re.I)
    v = re.sub(r"[_\-](out|pose\d+|model\d+)$", "", v, flags=re.I)
    return norm(v)


def fmt_num(value: Any, digits: int = 3) -> str:
    """Zahlen einheitlich formatieren, Nicht-Zahlen unveraendert lassen."""
    if value is None:
        return "-"
    s = str(value).strip()
    if s == "":
        return "-"
    try:
        f = float(s)
    except ValueError:
        return s
    if f != 0 and (abs(f) < 1e-3 or abs(f) >= 1e6):
        return f"{f:.3e}"
    return f"{f:.{digits}f}"


def breakable(s: str, chunk: int = 12) -> str:
    """Zero-Width-Spaces einfuegen, damit Word lange SMILES umbrechen kann."""
    if not s:
        return "-"
    return ZWSP.join(s[i:i + chunk] for i in range(0, len(s), chunk))


# ----------------------------------------------------------------------
# 0) MANIFEST.txt  (INI-artig, mit '#'-Kommentaren)
# ----------------------------------------------------------------------

def parse_ini_like(text: str) -> Dict[str, Dict[str, str]]:
    """
    Toleranter INI-Parser: ignoriert Freitext vor der ersten Sektion,
    schneidet '#'- und ';'-Kommentare ab, Keys werden kleingeschrieben.
    """
    data: Dict[str, Dict[str, str]] = {}
    section = "_GLOBAL"
    for raw in text.splitlines():
        line = raw.split("#", 1)[0].split(";", 1)[0].strip()
        if not line:
            continue
        if line.startswith("[") and line.endswith("]"):
            section = line[1:-1].strip().upper()
            data.setdefault(section, {})
            continue
        if "=" in line:
            k, v = line.split("=", 1)
            data.setdefault(section, {})[k.strip().lower()] = v.strip()
    return data


def mget(man: Dict[str, Dict[str, str]], key: str,
         default: str = "") -> str:
    """Wert ueber alle Sektionen hinweg suchen."""
    for sec in man.values():
        if key in sec:
            return sec[key]
    return default


def truthy(value: str) -> bool:
    return str(value).strip().lower() in ("true", "1", "yes", "on", "ja")


def as_float(value: str) -> Optional[float]:
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return None


def find_manifest(base: Path, csv_path: Optional[Path]) -> Optional[Path]:
    """MANIFEST.txt bevorzugt neben der Top250-CSV, sonst unter _unpacked."""
    roots: List[Path] = []
    if csv_path is not None:
        roots.append(csv_path.parent)
    unpacked = base / UNPACKED_DIRNAME
    if unpacked.is_dir():
        roots.append(unpacked)
    roots.append(base)

    for root in roots:
        for name in MANIFEST_NAMES:
            p = root / name
            if p.is_file():
                return p
    for root in roots:
        for name in MANIFEST_NAMES:
            hits = sorted(root.rglob(name))
            if hits:
                return hits[0]
    return None


def manifest_summary(path: Path) -> Dict[str, Any]:
    """Die fuer den Report relevanten Angaben aus dem MANIFEST ziehen."""
    man = parse_ini_like(path.read_text(errors="replace"))

    # --- aktive Gewichte (alles != 0) ---------------------------------
    weights: List[tuple[str, float]] = []
    for sec in man.values():
        for k, v in sec.items():
            if not k.startswith("w_"):
                continue
            f = as_float(v)
            if f is not None and f != 0.0:
                label = WEIGHT_LABELS.get(k, k)
                if all(label != w[0] for w in weights):
                    weights.append((label, f))
    weights.sort(key=lambda t: -t[1])
    wsum = sum(w for _, w in weights)

    # --- benutzte Modelle / Ensembles ---------------------------------
    models: List[tuple[str, str]] = []
    if truthy(mget(man, "dense_enabled")):
        models.append(("dense_model (CNNaffinity + CNNscore)",
                       mget(man, "dense_model", "-")))
    if truthy(mget(man, "cnnaffinity_enabled")) or \
       truthy(mget(man, "cnnscore_enabled")):
        models.append(("cnn_model", mget(man, "cnn_model", "-")))
    if truthy(mget(man, "deltalinf9xgb_enabled")):
        models.append(("Delta-Lin_F9XGB", "aktiv (eigenes Conda-Env)"))

    # --- aktive Rescoring-Terme ---------------------------------------
    resc_flags = [
        ("Vina", "vina_enabled"), ("Vinardo", "vinardo_enabled"),
        ("AD4", "ad4_enabled"), ("CNNaffinity", "cnnaffinity_enabled"),
        ("CNNscore", "cnnscore_enabled"), ("dense", "dense_enabled"),
        ("Delta-Lin_F9XGB", "deltalinf9xgb_enabled"),
    ]
    active_rescore = [lbl for lbl, key in resc_flags
                      if truthy(mget(man, key))]

    sigma = mget(man, "sigma_fraction", "-")
    clustering = "aus"
    if truthy(mget(man, "cluster_poses")):
        clustering = f"an, RMSD-Cutoff {mget(man, 'cluster_rmsd_cutoff', '?')} A"

    docking = [
        ("Engine / Binary", mget(man, "binary", "-")),
        ("Scoring-Funktion", mget(man, "scoring", "-")),
        ("Search-Mode", mget(man, "search_mode", "-")),
        ("Exhaustiveness", mget(man, "exhaustiveness", "-")),
        ("num_modes / energy_range",
         f"{mget(man, 'num_modes', '-')} / {mget(man, 'energy_range', '-')}"),
        ("Seed", (lambda s: "zufaellig (0)" if s.strip() == "0" else s or "-")(
            mget(man, "seed", ""))),
        ("GPU-Rescoring", "ja" if truthy(mget(man, "gnina_use_gpu")) else "nein"),
    ]

    return {"path": path, "raw": man, "weights": weights, "wsum": wsum,
            "models": models, "sigma": sigma, "clustering": clustering,
            "active_rescore": active_rescore, "docking": docking,
            "rescore_enabled": truthy(mget(man, "enabled", "true"))}


# ----------------------------------------------------------------------
# 1) SMILES aus PDBQT (nur MODEL 1)
# ----------------------------------------------------------------------

def extract_smiles(pdbqt: Path) -> str:
    """
    Liefert die SMILES-Notation aus dem ersten MODEL eines PDBQT.

    Meeko schreibt lange SMILES ueber mehrere 'REMARK SMILES'-Zeilen; diese
    werden ohne Trennzeichen zusammengehaengt. 'REMARK SMILES IDX' wird
    uebersprungen.
    """
    parts: List[str] = []
    try:
        with pdbqt.open("r", errors="replace") as fh:
            for line in fh:
                if line.startswith("ENDMDL"):
                    break                       # Ende von MODEL 1
                if not line.startswith("REMARK SMILES"):
                    continue
                rest = line[len("REMARK SMILES"):]
                if rest.lstrip().upper().startswith(("IDX", "INDEX")):
                    continue
                parts.append(rest.strip())
    except OSError as exc:
        print(f"  ! Lesefehler {pdbqt.name}: {exc}", file=sys.stderr)
        return ""
    return "".join(parts)


def load_rename_map(folder: Path) -> Dict[str, Dict[str, str]]:
    """rename_map.csv (von rename_pdbqt.py) einlesen -> {new_name: row}."""
    p = folder / RENAME_MAP_NAME
    if not p.is_file():
        return {}
    out: Dict[str, Dict[str, str]] = {}
    try:
        with p.open(newline="", errors="replace") as fh:
            for row in csv.DictReader(fh, delimiter=";"):
                if row.get("new_name"):
                    out[row["new_name"]] = row
    except OSError:
        return {}
    return out


def scan_folder(folder: Path) -> List[Dict[str, Any]]:
    """
    Alle Liganden-*.pdbqt eines Ordners einlesen -> Liste von Records.

    Achtung: bei der Benennung <rank>_<pocket>.pdbqt ist der Molekuelname
    NICHT mehr im Dateinamen enthalten - alle Dateien einer Serie wuerden
    denselben Namensschluessel ergeben. Die Records werden deshalb als Liste
    gefuehrt; build_table baut daraus die Indizes (Dateiname, Rank, Molekuel-
    name) und verwirft mehrdeutige Schluessel.

    selection.csv bzw. rename_map.csv liefern den urspruenglichen
    Ligandennamen, sofern vorhanden.
    """
    out: List[Dict[str, Any]] = []
    rmap = load_rename_map(folder)
    files = [f for f in sorted(folder.glob("*.pdbqt"))
             if f.name not in EXCLUDE_PDBQT
             and not f.name.lower().endswith(EXCLUDE_PDBQT_SUFFIX)]

    for f in files:
        m = FNAME_RE.match(f.name)
        rank = int(m.group("rank")) if m else None
        molname = m.group("name") if m else f.stem

        row = rmap.get(f.name)
        orig = row.get("old_name") if row else None
        if row and row.get("mol_name"):
            molname = row["mol_name"]

        out.append({
            "file_rank": rank,
            "mol_name": molname,
            "filename": f.name,
            "orig_name": orig,
            "smiles": extract_smiles(f),
        })

    n_sm = sum(1 for v in out if v["smiles"])
    note = f", rename_map.csv ({len(rmap)} Eintraege)" if rmap else ""
    print(f"  * {folder.name}: {len(files)} PDBQT gelesen "
          f"({n_sm} mit SMILES){note}")
    return out


# ----------------------------------------------------------------------
# 2) Top250-CSV
# ----------------------------------------------------------------------

def find_top_csv(base: Path, folder_name: str) -> Optional[Path]:
    """Top250_<folder>.csv rekursiv unter _unpacked suchen."""
    unpacked = base / UNPACKED_DIRNAME
    roots = [unpacked] if unpacked.is_dir() else [base]
    wanted = norm(f"Top250_{folder_name}")
    fallback: Optional[Path] = None
    for root in roots:
        for p in sorted(root.rglob(TOP_CSV_GLOB)):
            if norm(p.stem) == wanted:
                return p
            if norm(folder_name) in norm(p.stem) and fallback is None:
                fallback = p
    return fallback


def read_top_csv(path: Path) -> tuple[List[Dict[str, Any]], Dict[str, Optional[str]]]:
    """CSV einlesen, Delimiter sniffen, Spalten zuordnen."""
    with path.open("r", newline="", errors="replace") as fh:
        sample = fh.read(8192)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(fh, dialect=dialect)
        rows = [r for r in reader]
        fields = list(reader.fieldnames or [])

    cols = {k: find_col(fields, k) for k in COLUMN_CANDIDATES}
    return rows, cols


# ----------------------------------------------------------------------
# 3) Zusammenfuehren
# ----------------------------------------------------------------------

def find_score_csv(base: Path, folder: Path) -> Optional[Path]:
    """
    Score-Quelle bestimmen. Reihenfolge:
      1. selection.csv im Serienordner  (von extract_top_ligands.py)
      2. Top250_<serie>.csv unter _unpacked/
    """
    sel = folder / SELECTION_CSV_NAME
    if sel.is_file():
        return sel
    return find_top_csv(base, folder.name)


def build_table(folder: Path, base: Path, top_n: Optional[int]) -> Dict[str, Any]:
    records = scan_folder(folder)
    csv_path = find_score_csv(base, folder)

    rows: List[Dict[str, Any]] = []
    stats = {"csv_rows": 0, "matched": 0, "no_smiles": 0, "no_csv": 0}
    n_files = len(records)

    # Namensschluessel nur verwenden, wenn er eindeutig ist. Bei der Benennung
    # <rank>_<pocket>.pdbqt ergaeben alle Dateien denselben Schluessel.
    key_counts: Dict[str, int] = {}
    for rec in records:
        for cand in (rec["mol_name"], rec["orig_name"]):
            if cand:
                key_counts[strip_mol_key(cand)] = \
                    key_counts.get(strip_mol_key(cand), 0) + 1
    pdbqt: Dict[str, Dict[str, Any]] = {}
    for rec in records:
        for cand in (rec["mol_name"], rec["orig_name"]):
            if cand and key_counts.get(strip_mol_key(cand)) == 1:
                pdbqt[strip_mol_key(cand)] = rec

    if csv_path is None:
        print(f"  ! Keine Top250-CSV fuer {folder.name} gefunden "
              f"- nutze nur Dateinamen/SMILES", file=sys.stderr)
        cols = {k: None for k in COLUMN_CANDIDATES}
        for rec in sorted(records,
                          key=lambda r: (r["file_rank"] or 10**9)):
            rows.append({
                "rank": rec["file_rank"], "smiles": rec["smiles"],
                "ecr": "", "vina": "", "cnn": "", "pose": "",
                "mol_name": rec["mol_name"], "file": rec["filename"],
            })
    else:
        csv_rows, cols = read_top_csv(csv_path)
        stats["csv_rows"] = len(csv_rows)
        print(f"  * CSV: {csv_path.name}  ({len(csv_rows)} Zeilen)")
        print(f"    Spaltenzuordnung: " +
              ", ".join(f"{k}={v!r}" for k, v in cols.items() if v))

        # Zusatz-Indizes, falls der Name nicht matcht
        by_rank = {v["file_rank"]: v for v in records
                   if v["file_rank"] is not None}
        by_file: Dict[str, Dict[str, Any]] = {}
        for v in records:
            by_file[v["filename"]] = v
            if v.get("orig_name"):
                by_file[v["orig_name"]] = v

        for i, r in enumerate(csv_rows, start=1):
            raw_name = r.get(cols["name"], "") if cols["name"] else ""

            # 1. exakter Dateiname (selection.csv), 2. Molekuelname, 3. Rank
            rec = None
            if cols["file"]:
                rec = by_file.get(str(r.get(cols["file"], "")).strip())
            if rec is None:
                rec = pdbqt.get(strip_mol_key(raw_name))

            try:
                rank_val: Any = int(float(str(r.get(cols["rank"], i)).strip())) \
                    if cols["rank"] else i
            except (ValueError, TypeError):
                rank_val = r.get(cols["rank"], i) if cols["rank"] else i

            if rec is None and isinstance(rank_val, int):
                rec = by_rank.get(rank_val)

            smiles = (rec or {}).get("smiles", "")
            if cols["smiles"] and not smiles:
                smiles = r.get(cols["smiles"], "") or ""
            if rec is not None:
                stats["matched"] += 1
            if not smiles:
                stats["no_smiles"] += 1

            rows.append({
                "rank": rank_val,
                "smiles": smiles,
                "ecr": r.get(cols["ecr"], "") if cols["ecr"] else "",
                "vina": r.get(cols["vina"], "") if cols["vina"] else "",
                "cnn": r.get(cols["cnn"], "") if cols["cnn"] else "",
                "pose": r.get(cols["pose"], "") if cols["pose"] else "",
                "mol_name": (rec or {}).get("mol_name", raw_name),
                "file": (rec or {}).get("filename", ""),
            })

        matched_ranks = {r["rank"] for r in rows}
        stats["no_csv"] = sum(1 for v in records
                              if v["file_rank"] not in matched_ranks)

    rows.sort(key=lambda d: (d["rank"] if isinstance(d["rank"], int) else 10**9))
    if top_n:
        rows = rows[:top_n]

    renamed = any(v.get("orig_name") for v in records)

    # Dateinamensschema fuer den Hinweis im Report ableiten
    scheme = ""
    if records:
        first = min(records, key=lambda r: (r["file_rank"] or 10**9))
        m = FNAME_RE.match(first["filename"])
        scheme = (f"<RANK>_{m.group('name')}"
                  + ("_docked" if "_docked" in first["filename"] else "")
                  + ".pdbqt") if m else first["filename"]

    return {"folder": folder.name, "path": str(folder.resolve()),
            "csv": csv_path, "rows": rows, "stats": stats,
            "n_pdbqt": n_files, "renamed": renamed, "scheme": scheme}


# ----------------------------------------------------------------------
# 4) Word-Ausgabe
# ----------------------------------------------------------------------

def make_docx(datasets: List[Dict[str, Any]], base: Path, out: Path,
              title: str, subtitle: str,
              manifest: Optional[Dict[str, Any]] = None) -> None:
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError:
        sys.exit("FEHLER: python-docx fehlt.  ->  pip install python-docx")

    # ---- Layout-Konstanten (Template) --------------------------------
    ACCENT = RGBColor(0x1F, 0x3B, 0x57)      # Dunkelblau fuer Ueberschriften
    HEADER_BG = "1F3B57"                     # Tabellenkopf
    ZEBRA_BG = "F2F5F8"                      # jede 2. Zeile
    HINT_BG = "DCE6F1"                       # Hinweisbox mit Pfadangaben
    COL_W = [Cm(1.5), Cm(11.6), Cm(2.6), Cm(2.9), Cm(3.6), Cm(1.8)]
    HEADERS = ["RANK", "SMILES", "ECR Score", "score_vina_best",
               "score_dense_cnnaffinity_best", "Pose"]

    def shade(cell, hex_color: str) -> None:
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcpr.append(shd)

    def repeat_header(row) -> None:
        trpr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trpr.append(el)

    def put(cell, text: str, *, size: float, bold: bool = False,
            mono: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT,
            color: Optional[RGBColor] = None) -> None:
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = align
        pf = p.paragraph_format
        pf.space_before = Pt(1)
        pf.space_after = Pt(1)
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = "Consolas" if mono else "Calibri"
        if color is not None:
            run.font.color.rgb = color

    def page_number_footer(section) -> None:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("Seite ")
        run.font.size = Pt(8)
        for instr in ("PAGE", "NUMPAGES"):
            fld = OxmlElement("w:fldSimple")
            fld.set(qn("w:instr"), instr)
            p._p.append(fld)
            if instr == "PAGE":
                r = p.add_run(" von ")
                r.font.size = Pt(8)

    # ---- Dokument ----------------------------------------------------
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = (max(sec.page_width, sec.page_height),
                                       min(sec.page_width, sec.page_height))
    sec.left_margin = sec.right_margin = Cm(1.4)
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.4)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(9)

    sec.header.paragraphs[0].text = f"{title}  |  {base}"
    sec.header.paragraphs[0].runs[0].font.size = Pt(8)
    page_number_footer(sec)

    # ---- Titelblock --------------------------------------------------
    h = doc.add_heading(title, level=0)
    for r in h.runs:
        r.font.color.rgb = ACCENT
    p = doc.add_paragraph(subtitle)
    p.runs[0].font.italic = True

    meta = doc.add_table(rows=0, cols=2)
    meta.style = "Table Grid"
    total = sum(len(d["rows"]) for d in datasets)
    meta_rows = [
        ("Basisverzeichnis", str(base)),
        ("Erstellt am", _dt.datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Ausgewertete Serien", ", ".join(d["folder"] for d in datasets)),
        ("Eintraege gesamt", str(total)),
        ("SMILES-Quelle", "REMARK SMILES aus MODEL 1 der PDBQT-Dateien"),
        ("Score-Quelle", f"{SELECTION_CSV_NAME} im Serienordner, "
                         f"sonst Top250_<serie>.csv aus _unpacked/"),
        ("Posen-Ordner", ", ".join(f"{d['folder']}/" for d in datasets)
                        + "   (relativ zum Basisverzeichnis)"),
    ]
    if manifest:
        meta_rows.append(("Parameter-Quelle", str(manifest["path"])))

    for k, v in meta_rows:
        row = meta.add_row()
        put(row.cells[0], k, size=8.5, bold=True)
        put(row.cells[1], v, size=8.5)
        row.cells[0].width = Cm(5.0)
        row.cells[1].width = Cm(18.5)

    # ---- Parameter aus dem MANIFEST ----------------------------------
    def subheading(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = ACCENT

    def kv_table(pairs, w1=Cm(7.0), w2=Cm(16.5), head=None) -> None:
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        if head:
            hr = t.add_row()
            for i, (txt, w) in enumerate(zip(head, (w1, w2))):
                hr.cells[i].width = w
                shade(hr.cells[i], HEADER_BG)
                put(hr.cells[i], txt, size=8.5, bold=True,
                    color=RGBColor(0xFF, 0xFF, 0xFF))
        for i, (k, v) in enumerate(pairs):
            r = t.add_row()
            r.cells[0].width, r.cells[1].width = w1, w2
            put(r.cells[0], str(k), size=8.5, bold=True)
            put(r.cells[1], str(v), size=8.5)
            if i % 2 == 1:
                for c in r.cells:
                    shade(c, ZEBRA_BG)

    if manifest:
        subheading("Konsensus-Bewertung (ECR)")
        ecr_info = [
            ("sigma_fraction", manifest["sigma"]),
            ("Aktive Rescoring-Terme",
             ", ".join(manifest["active_rescore"]) or "-"),
            ("Pose-Clustering", manifest["clustering"]),
            ("Summe der Gewichte", fmt_num(manifest["wsum"], 2)),
        ]
        kv_table(ecr_info)

        subheading("Gewichtung der aktiven Scores")
        wpairs = [(lbl, f"{w:g}   ({w / manifest['wsum'] * 100:.0f} %)"
                   if manifest["wsum"] else f"{w:g}")
                  for lbl, w in manifest["weights"]]
        kv_table(wpairs or [("-", "keine Gewichte != 0 gefunden")],
                 head=("Score-Term", "Gewicht"))
        note = doc.add_paragraph(
            "Nicht aufgefuehrte Terme haben das Gewicht 0 und gehen nicht in "
            "den ECR-Score ein.")
        note.runs[0].font.size = Pt(8)
        note.runs[0].font.italic = True

        if manifest["models"]:
            subheading("Verwendete Modelle / Ensembles")
            kv_table(manifest["models"], head=("Modell", "Bezeichnung"))

        subheading("Docking-Parameter")
        kv_table(manifest["docking"])

    # ---- Je Serie eine Tabelle ---------------------------------------
    for ds in datasets:
        hh = doc.add_heading(f"Serie {ds['folder']}", level=1)
        hh.paragraph_format.page_break_before = True   # kein Leerseiten-Absatz
        for r in hh.runs:
            r.font.color.rgb = ACCENT

        src = ds["csv"].name if ds["csv"] else "keine CSV gefunden"
        info = doc.add_paragraph(
            f"PDBQT-Dateien: {ds['n_pdbqt']}   |   "
            f"Zeilen in Tabelle: {len(ds['rows'])}   |   "
            f"Score-Datei: {src}   |   "
            f"ohne SMILES: {ds['stats']['no_smiles']}"
            + ("   |   Dateien umbenannt (rename_map.csv)"
               if ds.get("renamed") else "")
        )
        info.runs[0].font.size = Pt(8)
        info.runs[0].font.italic = True

        # ---- Fundort der gedockten Posen -----------------------------
        box = doc.add_table(rows=0, cols=2)
        box.style = "Table Grid"
        box_rows = [("Posen-Verzeichnis", ds["path"] + "/")]
        if ds.get("scheme"):
            box_rows.append((
                "Dateischema",
                f"{ds['scheme']}   (<RANK> 4-stellig, entspricht Spalte RANK)"))
        hint = "Rezeptor: 00_target.pdbqt"
        if ds["csv"] and ds["csv"].name == SELECTION_CSV_NAME:
            hint += f"   |   Liganden-IDs je Rang: {SELECTION_CSV_NAME}"
        box_rows.append(("Im selben Ordner", hint))
        for i, (k, v) in enumerate(box_rows):
            row = box.add_row()
            row.cells[0].width, row.cells[1].width = Cm(4.4), Cm(19.1)
            shade(row.cells[0], HINT_BG)
            shade(row.cells[1], HINT_BG)
            put(row.cells[0], k, size=8, bold=True)
            put(row.cells[1], v, size=8, mono=True)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

        tbl = doc.add_table(rows=1, cols=6)
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        hdr = tbl.rows[0]
        repeat_header(hdr)
        for i, name in enumerate(HEADERS):
            c = hdr.cells[i]
            c.width = COL_W[i]
            shade(c, HEADER_BG)
            put(c, name, size=8.5, bold=True,
                align=WD_ALIGN_PARAGRAPH.CENTER,
                color=RGBColor(0xFF, 0xFF, 0xFF))

        for idx, r in enumerate(ds["rows"]):
            cells = tbl.add_row().cells
            for i in range(6):
                cells[i].width = COL_W[i]
            put(cells[0], str(r["rank"]), size=8, bold=True,
                align=WD_ALIGN_PARAGRAPH.CENTER)
            put(cells[1], breakable(r["smiles"]), size=6, mono=True)
            put(cells[2], fmt_num(r["ecr"], 4), size=8,
                align=WD_ALIGN_PARAGRAPH.RIGHT)
            put(cells[3], fmt_num(r["vina"], 3), size=8,
                align=WD_ALIGN_PARAGRAPH.RIGHT)
            put(cells[4], fmt_num(r["cnn"], 3), size=8,
                align=WD_ALIGN_PARAGRAPH.RIGHT)
            put(cells[5], str(r["pose"] or "-"), size=8,
                align=WD_ALIGN_PARAGRAPH.CENTER)
            if idx % 2 == 1:
                for c in cells:
                    shade(c, ZEBRA_BG)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"\n[OK] Word-Report geschrieben: {out}")


# ----------------------------------------------------------------------
# 5) Optionale Nebenausgaben
# ----------------------------------------------------------------------

def write_csv(datasets: List[Dict[str, Any]], out: Path) -> None:
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["SERIE", "RANK", "MOL_NAME", "DATEI", "PFAD", "SMILES",
                    "ECR", "score_vina_best",
                    "score_dense_cnnaffinity_best", "Pose"])
        for ds in datasets:
            for r in ds["rows"]:
                w.writerow([ds["folder"], r["rank"], r["mol_name"],
                            r.get("file", ""), ds["path"], r["smiles"],
                            r["ecr"], r["vina"], r["cnn"], r["pose"]])
    print(f"[OK] CSV geschrieben:          {out}")


def write_xlsx(datasets: List[Dict[str, Any]], out: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("  ! openpyxl fehlt - XLSX uebersprungen "
              "(pip install openpyxl)", file=sys.stderr)
        return
    wb = Workbook()
    wb.remove(wb.active)
    head = ["RANK", "SMILES", "ECR Score", "score_vina_best",
            "score_dense_cnnaffinity_best", "Pose", "MOL_NAME", "DATEI"]
    fill = PatternFill("solid", fgColor="1F3B57")
    for ds in datasets:
        ws = wb.create_sheet(ds["folder"][:31])
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = fill
        for r in ds["rows"]:
            ws.append([r["rank"], r["smiles"], r["ecr"], r["vina"],
                       r["cnn"], r["pose"], r["mol_name"],
                       r.get("file", "")])
        ws["A1"].comment = None
        for col, width in zip("ABCDEFGH", (8, 80, 14, 16, 26, 8, 34, 30)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        for c in ws["B"]:
            c.alignment = Alignment(wrap_text=False)
    wb.save(out)
    print(f"[OK] XLSX geschrieben:         {out}")


# ----------------------------------------------------------------------
# main
# ----------------------------------------------------------------------

def build_report(base: Path,
                 out: Optional[Path] = None,
                 *,
                 exclude: Optional[Sequence[str]] = None,
                 only: Optional[Sequence[str]] = None,
                 top: Optional[int] = None,
                 title: str = "Docking-Kaskade - Top-Hits Uebersicht",
                 subtitle: str = ("Konsensus-Ranking (ECR) mit AutoDock-Vina- "
                                  "und GNINA/CNN-Scores"),
                 manifest_path: Optional[Path] = None,
                 no_manifest: bool = False,
                 also_csv: bool = False,
                 also_xlsx: bool = False) -> Path:
    """
    Kompletten Report bauen. Wird sowohl von main() als auch von
    extract_top_ligands.py aufgerufen. Rueckgabe: Pfad der Word-Datei.
    """
    base = Path(base).expanduser().resolve()
    if not base.is_dir():
        sys.exit(f"FEHLER: Basisverzeichnis existiert nicht: {base}")

    excl = {e.strip() for e in (exclude if exclude is not None
                                else sorted(EXCLUDE_DIRS))}
    folders = sorted(p for p in base.iterdir()
                     if p.is_dir() and p.name not in excl
                     and not p.name.startswith("."))
    if only:
        folders = [f for f in folders if f.name in set(only)]
    if not folders:
        sys.exit(f"FEHLER: keine auswertbaren Ordner in {base}")

    print(f"Basis: {base}")
    print(f"Serien: {', '.join(f.name for f in folders)}\n")

    datasets = [build_table(f, base, top) for f in folders]

    # MANIFEST.txt fuer die Info-Seite
    manifest: Optional[Dict[str, Any]] = None
    if not no_manifest:
        mpath = manifest_path or find_manifest(
            base, next((d["csv"] for d in datasets if d["csv"]), None))
        if mpath and Path(mpath).is_file():
            try:
                manifest = manifest_summary(Path(mpath))
                print(f"\n  * MANIFEST: {mpath}")
                print(f"    sigma_fraction={manifest['sigma']}, "
                      f"aktive Gewichte: " +
                      ", ".join(f"{l}={w:g}" for l, w in manifest["weights"]))
            except Exception as exc:                       # noqa: BLE001
                print(f"  ! MANIFEST nicht auswertbar: {exc}", file=sys.stderr)
        else:
            print("  ! kein MANIFEST.txt gefunden - Parameter-Abschnitt "
                  "entfaellt", file=sys.stderr)

    out = Path(out) if out else (base / "docking_report.docx")
    make_docx(datasets, base, out, title, subtitle, manifest)

    if also_csv:
        write_csv(datasets, out.with_suffix(".csv"))
    if also_xlsx:
        write_xlsx(datasets, out.with_suffix(".xlsx"))
    return out


def main() -> None:
    ap = argparse.ArgumentParser(
        description="Docking-Auswertung (PDBQT + Top250-CSV) -> Word-Report",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    ap.add_argument("--base", type=Path, default=DEFAULT_BASE,
                    help="Basisverzeichnis mit den Serien-Ordnern")
    ap.add_argument("--out", type=Path, default=None,
                    help="Ziel-Word-Datei (Default: <base>/docking_report.docx)")
    ap.add_argument("--exclude", nargs="*", default=sorted(EXCLUDE_DIRS),
                    help="Ordnernamen, die uebersprungen werden")
    ap.add_argument("--only", nargs="*", default=None,
                    help="nur diese Serien auswerten")
    ap.add_argument("--top", type=int, default=None,
                    help="nur die besten N Eintraege je Serie")
    ap.add_argument("--title", default="Docking-Kaskade - Top-Hits Uebersicht")
    ap.add_argument("--subtitle",
                    default="Konsensus-Ranking (ECR) mit AutoDock-Vina- und "
                            "GNINA/CNN-Scores")
    ap.add_argument("--manifest", type=Path, default=None,
                    help="Pfad zur MANIFEST.txt (Default: Autosuche unter "
                         "_unpacked/)")
    ap.add_argument("--no-manifest", action="store_true",
                    help="Parameter-Abschnitt weglassen")
    ap.add_argument("--also-csv", action="store_true",
                    help="zusaetzlich eine Sammel-CSV schreiben")
    ap.add_argument("--also-xlsx", action="store_true",
                    help="zusaetzlich eine XLSX-Mappe schreiben")
    ap.add_argument("--dump-headers", action="store_true",
                    help="nur die Spaltennamen der Top250-CSVs anzeigen")
    args = ap.parse_args()

    base: Path = args.base.expanduser().resolve()
    if not base.is_dir():
        sys.exit(f"FEHLER: Basisverzeichnis existiert nicht: {base}")

    if args.dump_headers:
        excl = set(args.exclude)
        folders = sorted(p for p in base.iterdir()
                         if p.is_dir() and p.name not in excl
                         and not p.name.startswith("."))
        if args.only:
            folders = [f for f in folders if f.name in set(args.only)]
        for f in folders:
            p = find_score_csv(base, f)
            print(f"\n### {f.name}  ->  {p}")
            if p:
                _, cols = read_top_csv(p)
                with p.open(errors="replace") as fh:
                    print("Header:", fh.readline().rstrip())
                print("Zuordnung:", {k: v for k, v in cols.items() if v})
        return

    build_report(base, args.out, exclude=args.exclude, only=args.only,
                 top=args.top, title=args.title, subtitle=args.subtitle,
                 manifest_path=args.manifest, no_manifest=args.no_manifest,
                 also_csv=args.also_csv, also_xlsx=args.also_xlsx)


if __name__ == "__main__":
    main()
